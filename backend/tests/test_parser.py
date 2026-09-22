"""Parser tests: real files per format (PDF hand-built, no native deps)."""
from pathlib import Path

import pytest

from backend.app.services.ingestion.parser import SUPPORTED_EXTENSIONS, parse_file


def _make_pdf_bytes(text: str) -> bytes:
    """Minimal one-page PDF (avoids native PDF libs in test fixtures)."""
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode("latin-1") + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode("latin-1")
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode("latin-1")
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode("latin-1")
    return bytes(out)


@pytest.fixture
def sample_dir(tmp_path: Path) -> Path:
    (tmp_path / "leave.txt").write_text("Annual leave: 18 days per year.\nSick leave: 12 days.", encoding="utf-8")
    (tmp_path / "policy.md").write_text("# Leave Policy\n\nAnnual leave: 18 days.", encoding="utf-8")
    (tmp_path / "page.html").write_text("<html><body><h1>WFH</h1><p>2 days a week.</p></body></html>", encoding="utf-8")
    (tmp_path / "handbook.pdf").write_bytes(_make_pdf_bytes("Code of conduct: be respectful."))

    from docx import Document

    docx_path = tmp_path / "benefits.docx"
    d = Document()
    d.add_paragraph("Health insurance covers all full-time staff.")
    d.save(docx_path)

    from openpyxl import Workbook

    xlsx_path = tmp_path / "matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Notice"
    ws.append(["Tenure", "Notice"])
    ws.append(["0-2 years", "30 days"])
    wb.save(xlsx_path)
    return tmp_path


def test_supported_extensions():
    assert set(SUPPORTED_EXTENSIONS) >= {".pdf", ".docx", ".xlsx", ".txt", ".md", ".html"}


def test_parse_txt(sample_dir):
    pages = parse_file(sample_dir / "leave.txt")
    assert len(pages) == 1 and "18 days" in pages[0].text and pages[0].page_number == 1


def test_parse_md(sample_dir):
    pages = parse_file(sample_dir / "policy.md")
    assert "Leave Policy" in pages[0].text


def test_parse_html_strips_tags(sample_dir):
    pages = parse_file(sample_dir / "page.html")
    assert "<h1>" not in pages[0].text and "WFH" in pages[0].text


def test_parse_pdf(sample_dir):
    pages = parse_file(sample_dir / "handbook.pdf")
    assert len(pages) == 1 and "respectful" in pages[0].text and pages[0].page_number == 1


def test_parse_docx(sample_dir):
    pages = parse_file(sample_dir / "benefits.docx")
    assert "Health insurance" in pages[0].text


def test_parse_xlsx_sheet_as_page(sample_dir):
    pages = parse_file(sample_dir / "matrix.xlsx")
    assert len(pages) == 1 and "30 days" in pages[0].text and pages[0].section_heading == "Notice"


def test_unsupported_format(tmp_path):
    f = tmp_path / "a.csv"
    f.write_text("a,b", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_file(f)


def test_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        parse_file(tmp_path / "nope.pdf")
