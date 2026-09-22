"""Parse PDF/DOCX/XLSX/TXT/MD/HTML into raw text + page/section metadata.

Phase 1: digital documents only (no OCR — scanned PDFs need Tesseract, deferred).
`unstructured` lib intentionally skipped (heavy); direct parsers per format.
PDF uses pypdf (pure Python) — PyMuPDF's native DLL is blocked by this
machine's Application Control policy.
"""
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path


@dataclass
class ParsedPage:
    page_number: int
    text: str
    section_heading: str | None = None


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        stripped = data.strip()
        if stripped:
            self.parts.append(stripped)

    def get_text(self) -> str:
        return "\n".join(self.parts)


def _parse_pdf(path: Path) -> list[ParsedPage]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[ParsedPage] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(ParsedPage(page_number=i, text=text))
    return pages


def _parse_docx(path: Path) -> list[ParsedPage]:
    from docx import Document

    doc = Document(str(path))
    chunks: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                chunks.append(row_text)
    text = "\n".join(chunks).strip()
    return [ParsedPage(page_number=1, text=text)] if text else []


def _parse_xlsx(path: Path) -> list[ParsedPage]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages: list[ParsedPage] = []
    for i, sheet in enumerate(wb.worksheets, start=1):
        lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                lines.append(" | ".join(cells))
        if lines:
            pages.append(ParsedPage(page_number=i, text="\n".join(lines), section_heading=sheet.title))
    return pages


def _parse_text(path: Path) -> list[ParsedPage]:
    return [ParsedPage(page_number=1, text=path.read_text(encoding="utf-8").strip())]


def _parse_html(path: Path) -> list[ParsedPage]:
    extractor = _HTMLTextExtractor()
    extractor.feed(path.read_text(encoding="utf-8"))
    text = extractor.get_text().strip()
    return [ParsedPage(page_number=1, text=text)] if text else []


_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".txt": _parse_text,
    ".md": _parse_text,
    ".html": _parse_html,
    ".htm": _parse_html,
}

SUPPORTED_EXTENSIONS = sorted(_PARSERS.keys())


def parse_file(path: str | Path) -> list[ParsedPage]:
    """Parse a file into pages. Raises ValueError for unsupported formats, FileNotFoundError if missing."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"File not found: {p}")
    parser = _PARSERS.get(p.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported format '{p.suffix}'. Supported: {SUPPORTED_EXTENSIONS}")
    return parser(p)
